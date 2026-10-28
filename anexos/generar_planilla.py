#!/usr/bin/env python3
"""Genera anexos/modelo_financiero.xlsx: el modelo del anexo II con fórmulas vivas.

Es la misma cadena de cálculo que anexos/modelo_financiero.py, pero en una planilla que
recalcula al cambiar cualquier parámetro (celdas en azul). Los valores que anulan el VAN se
obtienen en forma cerrada: el VAN es lineal en el margen por sesión, así que no hace falta
Buscar objetivo.

Uso:  python3 anexos/generar_planilla.py   (requiere openpyxl)
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

SALIDA = Path(__file__).resolve().parent / "modelo_financiero.xlsx"

FUENTE = "Arial"
AZUL = Font(name=FUENTE, color="0000FF")
NEGRO = Font(name=FUENTE)
VERDE = Font(name=FUENTE, color="008000")
NEGRITA = Font(name=FUENTE, bold=True)
TITULO = Font(name=FUENTE, bold=True, size=14)
SUBTITULO = Font(name=FUENTE, bold=True, size=11, color="1F3864")
NOTA = Font(name=FUENTE, italic=True, color="595959", size=9)
ENCABEZADO = PatternFill("solid", fgColor="D9E2F3")
CLAVE = PatternFill("solid", fgColor="FFFF00")
LINEA = Border(bottom=Side(style="thin", color="8EA9DB"))

USD = '#,##0;(#,##0);"-"'
USD2 = '#,##0.00;(#,##0.00);"-"'
USD4 = '#,##0.0000;(#,##0.0000);"-"'
ARS = '"$"#,##0;("$"#,##0);"-"'
ARS2 = '"$"#,##0.00;("$"#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
ANIOS = '0.00 "años"'

ESCENARIOS = ["Pesimista", "Base", "Optimista"]
ALUMNOS = {
    "Pesimista": [130, 110, 160, 140, 190],
    "Base": [200, 260, 210, 290, 250],
    "Optimista": [280, 420, 340, 480, 400],
}

wb = Workbook()
nombres = {}


def nombrar(nombre, hoja, celda):
    """Nombre definido para una celda, así las fórmulas se leen como el modelo (=Ticket*Comision)."""
    col = "".join(c for c in celda if c.isalpha())
    fila = "".join(c for c in celda if c.isdigit())
    ref = f"'{hoja.title}'!${col}${fila}"
    wb.defined_names[nombre] = DefinedName(nombre, attr_text=ref)
    nombres[nombre] = ref


def escribir(hoja, celda, valor, fuente=NEGRO, formato=None, relleno=None):
    c = hoja[celda]
    c.value = valor
    c.font = fuente
    if formato:
        c.number_format = formato
    if relleno:
        c.fill = relleno
    return c


def encabezado(hoja, fila, textos, col=1):
    for i, t in enumerate(textos):
        c = hoja.cell(row=fila, column=col + i, value=t)
        c.font = NEGRITA
        c.fill = ENCABEZADO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def anchos(hoja, valores):
    for i, w in enumerate(valores, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = w


# ======================================================================= Léame
ley = wb.active
ley.title = "Léame"
anchos(ley, [110])
lineas = [
    ("Tinku — Modelo financiero (capítulo 5 y anexo II)", TITULO),
    ("", None),
    ("Planilla vigente del modelo con el que se calcularon los indicadores del capítulo 5. "
     "Reproduce las tablas del anexo II con fórmulas: al cambiar un parámetro, todo se recalcula.", NEGRO),
    ("Todos los montos están en dólares estadounidenses (USD) salvo que la columna diga ARS.", NEGRO),
    ("", None),
    ("Cómo se lee", SUBTITULO),
    ("Texto azul: parámetro de entrada (se puede modificar). Celdas amarillas: supuestos clave del modelo.", AZUL),
    ("Texto negro: fórmula calculada en la misma hoja. Texto verde: valor tomado de otra hoja.", NEGRO),
    ("", None),
    ("Hojas", SUBTITULO),
    ("Parámetros — supuestos del modelo con su origen (tabla de parámetros del anexo II).", NEGRO),
    ("Inversión — inversión inicial del Año 0 (tabla de inversión del capítulo 5).", NEGRO),
    ("Costos fijos — costos operativos mensuales y anuales (tabla de costos fijos del capítulo 5).", NEGRO),
    ("Demanda — alumnos activos y sesiones anuales por escenario (tabla de demanda del capítulo 5).", NEGRO),
    ("Fundador — horas y costo del tiempo del fundador en soporte, moderación y credenciales.", NEGRO),
    ("Flujos — flujos netos por escenario, sin y con el costo del fundador (y con el doble de horas), "
     "con VAN, TIR y repagos simple y descontado.", NEGRO),
    ("Indicadores — resumen de las tablas de indicadores del anexo II y del capítulo 5.", NEGRO),
    ("Sensibilidad — valores que anulan el VAN, comisión, tipo de cambio, ticket, duración, "
     "cota de inversión de USD 25.730 y riesgo R-09 (desarrollador rentado).", NEGRO),
    ("", None),
    ("Correspondencia", SUBTITULO),
    ("Genera esta planilla anexos/generar_planilla.py; anexos/modelo_financiero.py calcula los mismos "
     "valores en Python. Ambos están en el repositorio de la tesis (anexo I).", NEGRO),
    ("Los valores que anulan el VAN se calculan en forma cerrada (el VAN es lineal en el margen por sesión), "
     "por lo que no hace falta usar Buscar objetivo.", NEGRO),
]
for i, (texto, fuente) in enumerate(lineas, start=1):
    c = ley.cell(row=i, column=1, value=texto)
    c.font = fuente or NEGRO
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ================================================================== Parámetros
par = wb.create_sheet("Parámetros")
anchos(par, [46, 14, 14, 80])
escribir(par, "A1", "Parámetros del modelo", TITULO)
encabezado(par, 3, ["Parámetro", "Valor", "Unidad", "Origen"])
parametros = [
    # (nombre, etiqueta, valor o fórmula, unidad, formato, origen, clave)
    ("PrecioHoraARS", "Precio horario promedio", 13392.43, "ARS/hora", ARS2,
     "Relevamiento de 414 avisos de Superprof (apéndice E)", True),
    ("DuracionMin", "Duración media de la sesión", 51, "minutos", NUM,
     "Ecuación de duración media del capítulo 5", True),
    ("TipoCambio", "Tipo de cambio", 1535, "ARS/USD", NUM,
     "BNA, 09/09/2026 (cotización vendedor); el mismo que usa el sistema para el piso de tarifa", True),
    ("Ticket", "Ticket promedio por sesión", "=ROUND(PrecioHoraARS*DuracionMin/60/TipoCambio,2)", "USD/sesión",
     USD2, "Ecuación del ticket: precio horario × 51/60 / tipo de cambio, redondeado a centavos", False),
    ("Comision", "Comisión de la plataforma", 0.27, "%", PCT,
     "Regla BR-PAG-01 (ADR-M5-05)", True),
    ("ComisionMP", "Comisión de MercadoPago (a cargo del tutor)", 0.0604, "%", PCT2,
     "Modelo de cobro A (ADR-M5-02); no es costo de la plataforma, solo afecta el neto del tutor", False),
    ("Margen", "Margen de la plataforma por sesión", "=Ticket*Comision", "USD/sesión", USD4,
     "Ticket × comisión", False),
    ("ClasesMes", "Clases por alumno activo por mes", 3, "clases", NUM,
     "Supuesto del capítulo 5 (tabla de demanda)", True),
    ("Tasa", "Tasa de descuento", 0.15, "anual en USD", PCT,
     "Sección de tasa de descuento del capítulo 5", True),
    ("Inversion", "Inversión inicial (Año 0)", "='Inversión'!F13", "USD", USD,
     "Hoja Inversión", False),
    ("InversionCota", "Inversión de la cota (tarifas de contratista)", 25730, "USD", USD,
     "Versión anterior del capítulo 5: desarrollo valorizado a USD 50/30/15 por hora", False),
    ("PisoUSDHora", "Piso de tarifa por hora", 4, "USD/hora", USD2,
     "Piso del sistema (BR-PAG-04): $6.140 al tipo de cambio de $1.535", False),
]
fila = 4
for nombre, etiqueta, valor, unidad, formato, origen, clave in parametros:
    escribir(par, f"A{fila}", etiqueta)
    es_formula = isinstance(valor, str) and valor.startswith("=")
    fuente = (VERDE if "!" in valor else NEGRO) if es_formula else AZUL
    escribir(par, f"B{fila}", valor, fuente, formato, CLAVE if clave else None)
    escribir(par, f"C{fila}", unidad)
    escribir(par, f"D{fila}", origen, NOTA)
    nombrar(nombre, par, f"B{fila}")
    fila += 1

fila += 1
escribir(par, f"A{fila}", "Costo del tiempo del fundador", SUBTITULO)
fila += 1
encabezado(par, fila, ["Parámetro", "Valor", "Unidad", "Origen"])
fila += 1
fundador = [
    ("SalarioARS", "Salario mensual de referencia", 1316632, "ARS/mes", ARS,
     "Administrativo A, CCT 130/75, octubre de 2026, jornada completa (básico + suma fija); "
     "iProfesional, 01/10/2026", True),
    ("HorasMes", "Horas mensuales de jornada completa", "=48*52/12", "horas", NUM1,
     "48 horas semanales (Ley 11.544, art. 1) × 52 / 12", False),
    ("FactorSAC", "Factor de sueldo anual complementario", "=13/12", "", "0.0000",
     "Un sueldo más por año (Ley 20.744, art. 121)", False),
    ("ValorHoraUSD", "Valor hora del fundador", "=SalarioARS/HorasMes*FactorSAC/TipoCambio", "USD/hora",
     USD2, "Salario / horas × factor SAC / tipo de cambio", False),
    ("TasaConsulta", "Consultas de soporte por sesión", 0.05, "por sesión", "0.000",
     "Supuesto propio declarado en el capítulo 5", True),
    ("MinConsulta", "Minutos por consulta", 10, "min", NUM, "Supuesto propio", True),
    ("TasaDisputa", "Disputas o reembolsos manuales por sesión", 0.01, "por sesión", "0.000",
     "Supuesto propio declarado en el capítulo 5", True),
    ("MinDisputa", "Minutos por disputa", 20, "min", NUM, "Supuesto propio", True),
    ("TasaDenuncia", "Denuncias por sesión", 0.005, "por sesión", "0.000",
     "Supuesto propio declarado en el capítulo 5", True),
    ("MinDenuncia", "Minutos por denuncia", 45, "min", NUM, "Supuesto propio", True),
    ("AlumnosPorTutor", "Alumnos activos por tutor", 4, "alumnos", NUM,
     "Tutores activos = alumnos activos / 4", True),
    ("DocsPorTutor", "Documentos por tutor por año", 2, "documentos", NUM,
     "CAP (vigencia de 12 meses) + credencial académica", True),
    ("MinDocumento", "Minutos por documento revisado", 15, "min", NUM, "Supuesto propio", True),
    ("HorasFijasMes", "Horas fijas de revisión de paneles y colas", "=52/12", "horas/mes", NUM1,
     "1 hora semanal", False),
]
for nombre, etiqueta, valor, unidad, formato, origen, clave in fundador:
    escribir(par, f"A{fila}", etiqueta)
    es_formula = isinstance(valor, str) and valor.startswith("=")
    escribir(par, f"B{fila}", valor, NEGRO if es_formula else AZUL, formato, CLAVE if clave else None)
    escribir(par, f"C{fila}", unidad)
    escribir(par, f"D{fila}", origen, NOTA)
    nombrar(nombre, par, f"B{fila}")
    fila += 1

fila += 1
escribir(par, f"A{fila}", "Riesgo R-09: desarrollador rentado a jornada completa", SUBTITULO)
fila += 1
encabezado(par, fila, ["Parámetro", "Valor", "Unidad", "Origen"])
fila += 1
for nombre, etiqueta, valor, unidad, formato, origen in [
    ("HorasMesRentado", "Horas mensuales del desarrollador rentado", 160, "horas", NUM,
     "Mismo divisor que las tarifas de la tabla de inversión"),
]:
    escribir(par, f"A{fila}", etiqueta)
    escribir(par, f"B{fila}", valor, AZUL, formato)
    escribir(par, f"C{fila}", unidad)
    escribir(par, f"D{fila}", origen, NOTA)
    nombrar(nombre, par, f"B{fila}")
    fila += 1
for a in range(1, fila):
    par[f"D{a}"].alignment = Alignment(wrap_text=True, vertical="top")

# =================================================================== Inversión
inv = wb.create_sheet("Inversión")
anchos(inv, [16, 58, 12, 16, 14, 14])
escribir(inv, "A1", "Inversión inicial (Año 0)", TITULO)
encabezado(inv, 3, ["Categoría", "Concepto", "Cantidad", "Valor unitario (USD)", "Unidad", "Total (USD)"])
escribir(inv, "A15", "Tarifas por hora del personal", SUBTITULO)
encabezado(inv, 16, ["Perfil", "Mediana salarial bruta mensual (ARS)", "Horas/mes", "Tarifa (USD/hora)", "", ""])
for i, (perfil, salario) in enumerate([("Senior", 4100000), ("Semi-senior", 2800000), ("Junior", 1667500)]):
    r = 17 + i
    escribir(inv, f"A{r}", perfil)
    escribir(inv, f"B{r}", salario, AZUL, ARS)
    escribir(inv, f"C{r}", 160, AZUL, NUM)
    escribir(inv, f"D{r}", f"=ROUND(B{r}/C{r}/TipoCambio,0)", NEGRO, USD)
escribir(inv, "A20", "Medianas de desarrolladores con dedicación completa, Encuesta de Sueldos 2026.1 de Sysarmy; "
         "÷ 160 horas ÷ tipo de cambio, redondeadas a dólares enteros. No incluyen contribuciones patronales.", NOTA)
nombrar("TarifaSenior", inv, "D17")
nombrar("TarifaSemiSenior", inv, "D18")
nombrar("TarifaJunior", inv, "D19")

filas_inv = [
    ("Personal", "Senior — arquitectura, IA y WebRTC", 192, "=TarifaSenior", "horas"),
    ("Personal", "Semi-senior — backend y pasarela de pagos", 250, "=TarifaSemiSenior", "horas"),
    ("Personal", "Junior — frontend PWA e interfaz", 250, "=TarifaJunior", "horas"),
    ("Personal", "Junior — UAT y pruebas", 120, "=TarifaJunior", "horas"),
    ("Infraestructura", "Servidor propio preexistente y capas gratuitas (LiveKit, GitHub)", 6, 0, "meses"),
    ("Administrativo", "Registro de dominio", 1, 15, "año"),
    ("Legal", "Asesoría en privacidad de datos y TyC", 1, 800, "único"),
]
for i, (cat, concepto, cant, unit, unidad) in enumerate(filas_inv):
    r = 4 + i
    escribir(inv, f"A{r}", cat)
    escribir(inv, f"B{r}", concepto)
    escribir(inv, f"C{r}", cant, AZUL, NUM)
    escribir(inv, f"D{r}", unit, NEGRO if isinstance(unit, str) else AZUL, USD)
    escribir(inv, f"E{r}", unidad)
    escribir(inv, f"F{r}", f"=C{r}*D{r}", NEGRO, USD)
escribir(inv, "A11", "Riesgo")
escribir(inv, "B11", "Contingencia por imprevistos (% del desarrollo)")
escribir(inv, "C11", 0.10, AZUL, PCT)
escribir(inv, "D11", "=SUM(F4:F7)", NEGRO, USD)
escribir(inv, "E11", "base")
escribir(inv, "F11", "=ROUND(C11*D11,0)", NEGRO, USD)
escribir(inv, "B13", "Total inversión inicial", NEGRITA)
escribir(inv, "F13", "=SUM(F4:F10)+F11", NEGRITA, USD)
inv["F13"].border = LINEA

# ================================================================ Costos fijos
cf = wb.create_sheet("Costos fijos")
anchos(cf, [16, 40, 12, 12, 12, 12, 12])
escribir(cf, "A1", "Costos fijos operativos", TITULO)
encabezado(cf, 3, ["Categoría", "Concepto (USD por mes)", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
costos = [
    ("Infraestructura", "Servidor, base de datos y respaldos", [45, 45, 50, 55, 60]),
    ("Infraestructura", "Transmisión de video (LiveKit)", [60, 60, 65, 70, 75]),
    ("Administrativo", "Honorarios contables", [150, 150, 150, 150, 150]),
    ("Administrativo", "Renovación de dominio", [1.25] * 5),
    ("Comercial", "Adquisición de usuarios", [40, 50, 60, 70, 80]),
]
for i, (cat, concepto, valores) in enumerate(costos):
    r = 4 + i
    escribir(cf, f"A{r}", cat)
    escribir(cf, f"B{r}", concepto)
    for j, v in enumerate(valores):
        escribir(cf, f"{get_column_letter(3 + j)}{r}", v, AZUL, USD2)
escribir(cf, "B10", "Total mensual", NEGRITA)
escribir(cf, "B11", "Total anual", NEGRITA)
for j in range(5):
    col = get_column_letter(3 + j)
    escribir(cf, f"{col}10", f"=SUM({col}4:{col}8)", NEGRITA, USD2)
    escribir(cf, f"{col}11", f"={col}10*12", NEGRITA, USD)
    cf[f"{col}11"].border = LINEA
escribir(cf, "A13", "Fuente: tabla de costos fijos del capítulo 5.", NOTA)
COSTO_ANUAL = {t: f"'Costos fijos'!{get_column_letter(2 + t)}11" for t in range(1, 6)}

# ===================================================================== Demanda
dem = wb.create_sheet("Demanda")
anchos(dem, [30, 12, 12, 12, 12, 12, 14])
escribir(dem, "A1", "Demanda por escenario", TITULO)
encabezado(dem, 3, ["Alumnos activos", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5", "Promedio"])
encabezado(dem, 8, ["Sesiones anuales", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5", "Promedio"])
ALUM = {}
SES = {}
for i, esc in enumerate(ESCENARIOS):
    ra, rs = 4 + i, 9 + i
    escribir(dem, f"A{ra}", esc)
    escribir(dem, f"A{rs}", esc)
    for t, a in enumerate(ALUMNOS[esc], start=1):
        col = get_column_letter(1 + t)
        escribir(dem, f"{col}{ra}", a, AZUL, NUM, CLAVE)
        escribir(dem, f"{col}{rs}", f"={col}{ra}*ClasesMes*12", NEGRO, NUM)
        ALUM[(esc, t)] = f"Demanda!{col}{ra}"
        SES[(esc, t)] = f"Demanda!{col}{rs}"
    escribir(dem, f"G{ra}", f"=AVERAGE(B{ra}:F{ra})", NEGRO, NUM1)
    escribir(dem, f"G{rs}", f"=AVERAGE(B{rs}:F{rs})", NEGRO, NUM)
escribir(dem, "A13", "Sesiones anuales = alumnos activos × clases por mes × 12. Fuente: tabla de demanda del capítulo 5.",
         NOTA)

# ==================================================================== Fundador
fun = wb.create_sheet("Fundador")
anchos(fun, [34, 12, 12, 12, 12, 12])
escribir(fun, "A1", "Horas y costo del tiempo del fundador", TITULO)
escribir(fun, "A2", "horas/mes = sesiones/mes × (tasa × minutos)/60 por tarea + tutores × documentos/12 × minutos/60 "
         "+ horas fijas; costo anual = horas/mes × 12 × valor hora", NOTA)
COSTO_FUNDADOR = {}
fila = 4
for esc in ESCENARIOS:
    escribir(fun, f"A{fila}", f"Escenario {esc.lower()}", SUBTITULO)
    fila += 1
    encabezado(fun, fila, ["Concepto", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
    fila += 1
    conceptos = [
        ("Sesiones por mes", NUM),
        ("Tutores activos", NUM1),
        ("Horas de soporte", NUM1),
        ("Horas de soporte financiero", NUM1),
        ("Horas de moderación", NUM1),
        ("Horas de revisión de credenciales", NUM1),
        ("Horas fijas", NUM1),
        ("Total horas por mes", NUM1),
        ("Costo mensual (USD)", USD),
        ("Costo anual (USD)", USD),
    ]
    base_fila = fila
    for k, (etiqueta, _) in enumerate(conceptos):
        escribir(fun, f"A{fila + k}", etiqueta, NEGRITA if etiqueta.startswith(("Total", "Costo")) else NEGRO)
    for t in range(1, 6):
        col = get_column_letter(1 + t)
        s, tu = f"{col}{base_fila}", f"{col}{base_fila + 1}"
        formulas = [
            f"={ALUM[(esc, t)]}*ClasesMes",
            f"={ALUM[(esc, t)]}/AlumnosPorTutor",
            f"={s}*TasaConsulta*MinConsulta/60",
            f"={s}*TasaDisputa*MinDisputa/60",
            f"={s}*TasaDenuncia*MinDenuncia/60",
            f"={tu}*DocsPorTutor/12*MinDocumento/60",
            "=HorasFijasMes",
            f"=SUM({col}{base_fila + 2}:{col}{base_fila + 6})",
            f"={col}{base_fila + 7}*ValorHoraUSD",
            f"={col}{base_fila + 8}*12",
        ]
        for k, f in enumerate(formulas):
            escribir(fun, f"{col}{base_fila + k}", f, NEGRO, conceptos[k][1])
        COSTO_FUNDADOR[(esc, t)] = f"Fundador!{col}{base_fila + 9}"
    fila = base_fila + len(conceptos) + 1

# ================================================================== Flujos
flu = wb.create_sheet("Flujos")
anchos(flu, [40, 13, 13, 13, 13, 13, 13])
escribir(flu, "A1", "Flujos netos e indicadores por escenario (USD)", TITULO)
escribir(flu, "A2", "VAN = Σ flujo_t / (1 + tasa)^t; TIR con IRR; repago = año previo al cruce + saldo pendiente "
         "/ flujo del año del cruce (simple sobre flujos, descontado sobre flujos descontados).", NOTA)
INDIC = {}


def bloque_indicadores(hoja, fila, etiqueta, celdas_flujo, clave):
    """Escribe flujo, acumulados y repagos para una serie de 6 celdas (Año 0 a 5).

    celdas_flujo: lista de 6 fórmulas o referencias para la fila de flujos.
    Devuelve la fila siguiente libre y guarda las referencias de los indicadores en INDIC[clave].
    """
    rf, ra, rr, rd, rad, rrd = range(fila, fila + 6)
    etiquetas = [etiqueta, "  acumulado", "  cruce (repago simple)", "  flujo descontado",
                 "  acumulado descontado", "  cruce (repago descontado)"]
    for k, e in enumerate(etiquetas):
        escribir(hoja, f"A{fila + k}", e, NEGRITA if k == 0 else NOTA)
    for t in range(6):
        col = get_column_letter(2 + t)
        prev = get_column_letter(1 + t)
        escribir(hoja, f"{col}{rf}", celdas_flujo[t], NEGRITA, USD)
        escribir(hoja, f"{col}{ra}", f"={col}{rf}" if t == 0 else f"={prev}{ra}+{col}{rf}", NEGRO, USD)
        escribir(hoja, f"{col}{rd}", f"={col}{rf}/(1+Tasa)^{t}", NEGRO, USD)
        escribir(hoja, f"{col}{rad}", f"={col}{rd}" if t == 0 else f"={prev}{rad}+{col}{rd}", NEGRO, USD)
        if t > 0:
            escribir(hoja, f"{col}{rr}",
                     f'=IF(AND({prev}{ra}<0,{col}{ra}>=0),{t - 1}+ABS({prev}{ra})/{col}{rf},"")', NEGRO, "0.00")
            escribir(hoja, f"{col}{rrd}",
                     f'=IF(AND({prev}{rad}<0,{col}{rad}>=0),{t - 1}+ABS({prev}{rad})/{col}{rd},"")', NEGRO, "0.00")
    rango = f"B{rf}:G{rf}"
    # Con flujos que cambian de signo varias veces la TIR tiene más de una raíz y Newton (IRR) puede
    # caer en una menor que -100 %, sin sentido económico: en ese caso se reintenta desde -50 %.
    INDIC[clave] = {
        "flujos": [f"'{hoja.title}'!{get_column_letter(2 + t)}{rf}" for t in range(6)],
        "van": f"=NPV(Tasa,'{hoja.title}'!C{rf}:G{rf})+'{hoja.title}'!B{rf}",
        "tir": f"=IFERROR(IF(IRR('{hoja.title}'!{rango},0)<-0.99,IRR('{hoja.title}'!{rango},-0.5),IRR('{hoja.title}'!{rango},0)),\"no existe\")",
        "repago": f"=IF(COUNT('{hoja.title}'!C{rr}:G{rr})=0,\"no\",MIN('{hoja.title}'!C{rr}:G{rr}))",
        "repago_desc": f"=IF(COUNT('{hoja.title}'!C{rrd}:G{rrd})=0,\"no\",MIN('{hoja.title}'!C{rrd}:G{rrd}))",
    }
    return fila + 7


fila = 4
for esc in ESCENARIOS:
    escribir(flu, f"A{fila}", f"Escenario {esc.lower()}", SUBTITULO)
    fila += 1
    encabezado(flu, fila, ["Concepto", "Año 0", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
    fila += 1
    r_ses, r_ing, r_cf, r_fun = fila, fila + 1, fila + 2, fila + 3
    escribir(flu, f"A{r_ses}", "Sesiones anuales")
    escribir(flu, f"A{r_ing}", "Ingreso neto (sesiones × margen)")
    escribir(flu, f"A{r_cf}", "Costos fijos anuales")
    escribir(flu, f"A{r_fun}", "Costo del tiempo del fundador")
    for t in range(1, 6):
        col = get_column_letter(2 + t)
        escribir(flu, f"{col}{r_ses}", f"={SES[(esc, t)]}", VERDE, NUM)
        escribir(flu, f"{col}{r_ing}", f"={col}{r_ses}*Margen", NEGRO, USD)
        escribir(flu, f"{col}{r_cf}", f"={COSTO_ANUAL[t]}", VERDE, USD)
        escribir(flu, f"{col}{r_fun}", f"={COSTO_FUNDADOR[(esc, t)]}", VERDE, USD)
    fila += 5
    for variante, factor in (("sin", 0), ("con", 1), ("doble", 2)):
        etiqueta = {"sin": "Flujo neto sin costo del fundador",
                    "con": "Flujo neto con costo del fundador",
                    "doble": "Flujo neto con el doble de horas"}[variante]
        celdas = ["=-Inversion"] + [
            f"={get_column_letter(2 + t)}{r_ing}-{get_column_letter(2 + t)}{r_cf}"
            + ("" if factor == 0 else f"-{factor}*{get_column_letter(2 + t)}{r_fun}")
            for t in range(1, 6)
        ]
        fila = bloque_indicadores(flu, fila, etiqueta, celdas, (esc, variante))
    fila += 1

# ================================================================= Indicadores
ind = wb.create_sheet("Indicadores")
anchos(ind, [14, 28, 14, 14, 16, 18])
escribir(ind, "A1", "Indicadores por escenario", TITULO)
encabezado(ind, 3, ["Escenario", "Variante", "VAN (USD)", "TIR", "Repago simple", "Repago descontado"])
fila = 4
for esc in ESCENARIOS:
    for variante, etiqueta in (("sin", "Sin costo del fundador"), ("con", "Con costo del fundador"),
                               ("doble", "Con el doble de horas")):
        d = INDIC[(esc, variante)]
        escribir(ind, f"A{fila}", esc if variante == "sin" else "")
        escribir(ind, f"B{fila}", etiqueta)
        escribir(ind, f"C{fila}", d["van"], VERDE, USD)
        escribir(ind, f"D{fila}", d["tir"], VERDE, PCT2)
        escribir(ind, f"E{fila}", d["repago"], VERDE, ANIOS)
        escribir(ind, f"F{fila}", d["repago_desc"], VERDE, ANIOS)
        fila += 1
escribir(ind, f"A{fila + 1}", "Las filas «sin costo del fundador» corresponden a la tabla de indicadores del anexo II; "
         "las tres variantes, a la tabla del VAN con y sin el costo del fundador.", NOTA)

# ================================================================ Sensibilidad
sen = wb.create_sheet("Sensibilidad")
anchos(sen, [44, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14])
escribir(sen, "A1", "Análisis de sensibilidad (escenario base salvo indicación)", TITULO)
escribir(sen, "A2", "El VAN es lineal en el margen m: VAN(m) = m × VA(sesiones) − inversión − VA(costos fijos). "
         "Los valores que lo anulan se despejan de esa ecuación.", NOTA)
encabezado(sen, 4, ["Valor presente (Años 1 a 5)", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5", "Suma"])
escribir(sen, "A5", "Factor de descuento")
escribir(sen, "A6", "Costos fijos anuales")
escribir(sen, "A7", "Sesiones anuales, base")
escribir(sen, "A8", "Sesiones anuales, pesimista")
for t in range(1, 6):
    col = get_column_letter(1 + t)
    escribir(sen, f"{col}5", f"=1/(1+Tasa)^{t}", NEGRO, "0.0000")
    escribir(sen, f"{col}6", f"={COSTO_ANUAL[t]}", VERDE, USD)
    escribir(sen, f"{col}7", f"={SES[('Base', t)]}", VERDE, NUM)
    escribir(sen, f"{col}8", f"={SES[('Pesimista', t)]}", VERDE, NUM)
escribir(sen, "G5", "=SUM(B5:F5)", NEGRO, "0.0000")
escribir(sen, "G6", "=SUMPRODUCT(B6:F6,B5:F5)", NEGRO, USD)
escribir(sen, "G7", "=SUMPRODUCT(B7:F7,B5:F5)", NEGRO, NUM)
escribir(sen, "G8", "=SUMPRODUCT(B8:F8,B5:F5)", NEGRO, NUM)
nombrar("AnualidadVA", sen, "G5")
nombrar("VACostos", sen, "G6")
nombrar("VASesBase", sen, "G7")
nombrar("VASesPes", sen, "G8")

encabezado(sen, 10, ["Valor que anula el VAN", "Valor", "Unidad", "", "", "", ""])
anulan = [
    ("Margen por sesión, base", "=(Inversion+VACostos)/VASesBase", "USD/sesión", USD4, "MargenEq"),
    ("Comisión, base", "=MargenEq/Ticket", "%", PCT, None),
    ("Comisión, pesimista", "=(Inversion+VACostos)/VASesPes/Ticket", "%", PCT, None),
    ("Ticket por sesión", "=MargenEq/Comision", "USD/sesión", USD2, "TicketEq"),
    ("Ticket por hora", "=TicketEq*60/DuracionMin", "USD/hora", USD2, None),
    ("Ticket por hora en pesos", "=TicketEq*60/DuracionMin*TipoCambio", "ARS/hora", ARS, None),
    ("Tipo de cambio", "=PrecioHoraARS*DuracionMin/60/TicketEq", "ARS/USD", NUM, None),
    ("Caída del volumen de sesiones", "=1-MargenEq/Margen", "%", PCT, None),
    ("Sesiones anuales promedio", "=MargenEq/Margen*Demanda!G10", "sesiones", NUM, None),
    ("Descuento medio de paquete", "=1-TicketEq/Ticket", "%", PCT, None),
    ("Duración media de la sesión", "=DuracionMin*MargenEq/Margen", "minutos", NUM1, None),
    ("Costo mensual constante adicional", f"=({INDIC[('Base', 'sin')]['van'][1:]})/AnualidadVA/12", "USD/mes", USD,
     "ExcedenteMes"),
    ("  equivalente en horas del fundador", "=ExcedenteMes/ValorHoraUSD", "horas/mes", NUM, None),
    ("Inversión inicial, pesimista", f"={INDIC[('Pesimista', 'sin')]['van'][1:]}+Inversion", "USD", USD, None),
    ("Inversión inicial, base", f"={INDIC[('Base', 'sin')]['van'][1:]}+Inversion", "USD", USD, None),
    ("Inversión inicial, optimista", f"={INDIC[('Optimista', 'sin')]['van'][1:]}+Inversion", "USD", USD, None),
]
fila = 11
for etiqueta, formula, unidad, formato, nombre in anulan:
    escribir(sen, f"A{fila}", etiqueta)
    escribir(sen, f"B{fila}", formula, NEGRO, formato)
    escribir(sen, f"C{fila}", unidad)
    if nombre:
        nombrar(nombre, sen, f"B{fila}")
    fila += 1

fila += 1
encabezado(sen, fila, ["Ticket alternativo (escenario base)", "Ticket (USD)", "VAN base (USD)", "", "", "", ""])
fila += 1
for etiqueta, ticket in [
    ("Todas las sesiones al piso de tarifa", "=PisoUSDHora*DuracionMin/60"),
    ("Mitad de las sesiones en paquete con 30 %", "=Ticket*0.85"),
    ("Todas las sesiones en paquete con 30 %", "=Ticket*0.7"),
    ("Duración de 45 minutos", "=ROUND(PrecioHoraARS/TipoCambio*45/60,2)"),
    ("Duración de 60 minutos", "=ROUND(PrecioHoraARS/TipoCambio*60/60,2)"),
]:
    escribir(sen, f"A{fila}", etiqueta)
    escribir(sen, f"B{fila}", ticket, NEGRO, USD2)
    escribir(sen, f"C{fila}", f"=B{fila}*Comision*VASesBase-Inversion-VACostos", NEGRO, USD)
    fila += 1

fila += 1
escribir(sen, f"A{fila}", "Comisión de la plataforma", SUBTITULO)
fila += 1
encabezado(sen, fila, ["Comisión", "Margen", "Año 0", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5",
                       "VAN base", "TIR base", "VAN pesimista"])
sen.cell(row=fila, column=12, value="Neto del tutor").font = NEGRITA
sen.cell(row=fila, column=12).fill = ENCABEZADO
fila += 1
for c in ["=MargenEq/Ticket", 0.21, 0.23, 0.25, "=Comision", 0.30, 0.33]:
    escribir(sen, f"A{fila}", c, NEGRO if isinstance(c, str) else AZUL, PCT)
    escribir(sen, f"B{fila}", f"=Ticket*A{fila}", NEGRO, USD4)
    escribir(sen, f"C{fila}", "=-Inversion", NEGRO, USD)
    for t in range(1, 6):
        escribir(sen, f"{get_column_letter(3 + t)}{fila}", f"={SES[('Base', t)]}*B{fila}-{COSTO_ANUAL[t]}", NEGRO, USD)
    escribir(sen, f"I{fila}", f"=NPV(Tasa,D{fila}:H{fila})+C{fila}", NEGRO, USD)
    escribir(sen, f"J{fila}", f'=IFERROR(IF(IRR(C{fila}:H{fila},0)<-0.99,IRR(C{fila}:H{fila},-0.5),IRR(C{fila}:H{fila},0)),"no existe")', NEGRO, PCT2)
    escribir(sen, f"K{fila}", f"=B{fila}*VASesPes-Inversion-VACostos", NEGRO, USD)
    escribir(sen, f"L{fila}", f"=1-A{fila}-ComisionMP", NEGRO, PCT2)
    fila += 1
sen.column_dimensions["L"].width = 14

fila += 1
escribir(sen, f"A{fila}", "Tipo de cambio", SUBTITULO)
fila += 1
encabezado(sen, fila, ["Tipo de cambio (ARS/USD)", "Ticket", "Margen", "Año 0", "Año 1", "Año 2", "Año 3",
                       "Año 4", "Año 5", "VAN base", "TIR base"])
for col, t in (("L", "Sesiones eq. Año 1"), ("M", "Sesiones eq. Año 5")):
    sen[f"{col}{fila}"] = t
    sen[f"{col}{fila}"].font = NEGRITA
    sen[f"{col}{fila}"].fill = ENCABEZADO
    sen.column_dimensions[col].width = 16
fila += 1
for tc in [1400, "=TipoCambio", 1700, 1900, 2100]:
    escribir(sen, f"A{fila}", tc, NEGRO if isinstance(tc, str) else AZUL, NUM)
    escribir(sen, f"B{fila}", f"=IF(A{fila}=TipoCambio,Ticket,ROUND(PrecioHoraARS*DuracionMin/60/A{fila},2))",
             NEGRO, USD2)
    escribir(sen, f"C{fila}", f"=B{fila}*Comision", NEGRO, USD2)
    escribir(sen, f"D{fila}", "=-Inversion", NEGRO, USD)
    for t in range(1, 6):
        escribir(sen, f"{get_column_letter(4 + t)}{fila}", f"={SES[('Base', t)]}*C{fila}-{COSTO_ANUAL[t]}", NEGRO, USD)
    escribir(sen, f"J{fila}", f"=NPV(Tasa,E{fila}:I{fila})+D{fila}", NEGRO, USD)
    escribir(sen, f"K{fila}", f'=IFERROR(IF(IRR(D{fila}:I{fila},0)<-0.99,IRR(D{fila}:I{fila},-0.5),IRR(D{fila}:I{fila},0)),"no existe")', NEGRO, PCT2)
    escribir(sen, f"L{fila}", f"={COSTO_ANUAL[1]}/C{fila}", NEGRO, NUM)
    escribir(sen, f"M{fila}", f"={COSTO_ANUAL[5]}/C{fila}", NEGRO, NUM)
    fila += 1

fila += 1
escribir(sen, f"A{fila}", "Cota: inversión valorizada a tarifas de contratista (sin costo del fundador)", SUBTITULO)
fila += 2
cota_inicio = fila
for esc in ESCENARIOS:
    celdas = ["=-InversionCota"] + [f"={INDIC[(esc, 'sin')]['flujos'][t]}" for t in range(1, 6)]
    fila = bloque_indicadores(sen, fila, f"Flujo neto {esc.lower()}", celdas, (esc, "cota"))
encabezado(sen, cota_inicio - 1, ["Concepto", "Año 0", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
encabezado(sen, fila, ["Escenario", "VAN (USD)", "TIR", "Repago simple", "Repago descontado", "", ""])
fila += 1
for esc in ESCENARIOS:
    d = INDIC[(esc, "cota")]
    escribir(sen, f"A{fila}", esc)
    escribir(sen, f"B{fila}", d["van"], NEGRO, USD)
    escribir(sen, f"C{fila}", d["tir"], NEGRO, PCT2)
    escribir(sen, f"D{fila}", d["repago"], NEGRO, ANIOS)
    escribir(sen, f"E{fila}", d["repago_desc"], NEGRO, ANIOS)
    fila += 1

fila += 1
escribir(sen, f"A{fila}", "Riesgo R-09: desarrollador rentado a jornada completa (sin costo del fundador)", SUBTITULO)
fila += 1
encabezado(sen, fila, ["Perfil / escenario", "Costo anual", "Año 0", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5",
                       "VAN", "TIR", ""])
fila += 1
for perfil, tarifa in (("Junior", "TarifaJunior"), ("Semi-senior", "TarifaSemiSenior")):
    escribir(sen, f"A{fila}", perfil, NEGRITA)
    escribir(sen, f"B{fila}", f"={tarifa}*HorasMesRentado*12", NEGRO, USD)
    r_perfil = fila
    fila += 1
    for esc in ESCENARIOS:
        escribir(sen, f"A{fila}", f"  {esc}")
        escribir(sen, f"C{fila}", f"={INDIC[(esc, 'sin')]['flujos'][0]}", VERDE, USD)
        for t in range(1, 6):
            escribir(sen, f"{get_column_letter(3 + t)}{fila}",
                     f"={INDIC[(esc, 'sin')]['flujos'][t]}-$B${r_perfil}", NEGRO, USD)
        escribir(sen, f"I{fila}", f"=NPV(Tasa,D{fila}:H{fila})+C{fila}", NEGRO, USD)
        escribir(sen, f"J{fila}", f'=IFERROR(IF(IRR(C{fila}:H{fila},0)<-0.99,IRR(C{fila}:H{fila},-0.5),IRR(C{fila}:H{fila},0)),"no existe")', NEGRO, PCT2)
        fila += 1

# ------------------------------------------------------------------ cierre
for hoja in wb.worksheets:
    hoja.sheet_view.showGridLines = False
    # Impresión: una página de ancho, en horizontal.
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    if hoja.title != "Léame":
        hoja.freeze_panes = "B4" if hoja.title not in ("Parámetros",) else "A4"
wb.save(SALIDA)
print(f"Escrito {SALIDA}")
